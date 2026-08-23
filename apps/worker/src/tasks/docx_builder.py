"""Builds the SRS document directly via python-docx — deliberately NOT
docxtpl for the body content.

Every text insertion goes through python-docx's own `.text =` / `add_run()`
APIs, which XML-escape content correctly by construction (they operate on
lxml element objects, not raw XML strings) — this is what actually
eliminates the class of document-corruption bug docxtpl's raw-string Jinja
templating was exposed to, not just a patch on top of it.

The base document is the org's real template
(`templates/main_template_SRS-Customer-BankAsiaSmartApp-V0.5.8.docx`), not a
blank `Document()` — loading it and stripping only the body content (while
keeping the trailing `sectPr`) carries over its header/footer (ERA logo on
every page, "CONFIDENTIAL" banner, copyright footer), fonts, theme, and
named paragraph/table styles for free, without ever touching XML as a raw
string. This is the same "template as base document" technique used
everywhere else in this pipeline in place of docxtpl.

Renders the real Epic > Feature > Story hierarchy (Tasks/Bugs as a
checklist under their parent, not their own sections), acceptance
criteria, every embedded image, a Word Table of Contents field, and a
traceability matrix as a native table.
"""

from __future__ import annotations

import logging
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image
from srs_core.rendering.html_text import blocks_to_plain_text
from srs_core.security.file_signature import sniff_content_type

logger = logging.getLogger(__name__)

_LEAF_TYPES = {"task", "bug"}

# The org template's real paragraph styles (verified against
# templates/main_template_SRS-Customer-BankAsiaSmartApp-V0.5.8.docx):
# "S Chapter Focus" for the cover title, "S Label" for cover metadata lines,
# "S Heading"/"S Heading 2"/"S Heading 3" for Epic/Feature/Story headings
# respectively (the template only exercises "S Heading" once, at Epic level
# — Feature headings, e.g. "Feature-001: Authentication", and top-level
# document sections like "Introduction" both use "S Heading 2"), "S Notes"
# for narrative paragraphs, "List Bullet" for bullets.
STYLE_COVER_TITLE = "S Chapter Focus"
STYLE_COVER_LABEL = "S Label"
STYLE_SECTION_HEADING = "S Heading 2"
STYLE_EPIC_HEADING = "S Heading"
STYLE_FEATURE_HEADING = "S Heading 2"
STYLE_STORY_HEADING = "S Heading 3"
STYLE_NARRATIVE = "S Notes"
# "List Paragraph" is what the template's SAMPLE content used visually, but
# it carries no numbering of its own — a paragraph in that style renders as
# plain, unmarked text with no bullet glyph at all (confirmed against the
# template's styles.xml: no <w:numPr> anywhere in its definition). "List
# Bullet" is the style Word itself defines WITH numbering baked in.
STYLE_BULLET = "List Bullet"
STYLE_NUMBER = "List Number"  # for <ol>-derived Markdown lists — falls back gracefully if the template lacks it
TABLE_STYLE = "Table Grid"
_CODE_FONT_NAME = "Consolas"
_CODE_BLOCK_FONT_SIZE_PT = 9.5
_CODE_BLOCK_SHADING_FILL = "D9D9D9"  # light grey

_TEMPLATE_DIR = Path(__file__).resolve().parents[4] / "templates"
ORG_TEMPLATE_PATH = _TEMPLATE_DIR / "main_template_SRS-Customer-BankAsiaSmartApp-V0.5.8.docx"
_HEADER_APP_NAME_PLACEHOLDER = "Mobile Banking App"
_FOOTER_YEAR_PLACEHOLDER = "2025"


def _safe_add_paragraph(document_or_cell, text: str = "", style: str | None = None):
    """`document.add_paragraph(text, style=name)` raises KeyError if the
    named style doesn't exist in this document's styles part — falls back to
    the default style rather than failing the whole render, in case a future
    template swap doesn't define one of these exact style names.

    Defaults every paragraph to LEFT alignment regardless of what its named
    style specifies — the org template's "S Notes" style (used for most body
    content) defaults to JUSTIFY, which is fine for full-width prose but
    stretches any short line (a table row, a single bullet, a short
    heading) into unreadable, widely-spaced text. Callers that genuinely
    want centered text (the cover page) explicitly set `.alignment` on the
    returned paragraph afterward, which overrides this default.
    """
    try:
        paragraph = document_or_cell.add_paragraph(text, style=style) if style else document_or_cell.add_paragraph(text)
    except KeyError:
        logger.warning("style %r not found in template, falling back to default", style)
        paragraph = document_or_cell.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return paragraph


def _apply_run_formatting(run, run_data: dict) -> None:
    if run_data.get("bold"):
        run.bold = True
    if run_data.get("italic"):
        run.italic = True
    if run_data.get("underline"):
        run.underline = True
    if run_data.get("mono"):
        run.font.name = _CODE_FONT_NAME


def _add_runs_paragraph(document_or_cell, runs: list[dict], style: str | None = None):
    """Renders an `html_to_blocks` run list (from a `text` block or a list
    item) as one paragraph with per-run bold/italic/underline preserved —
    inline formatting like a bolded "Short CIF"/"Full CIF" inside an
    otherwise plain sentence used to be silently discarded when content was
    flattened to a single plain string.
    """
    paragraph = _safe_add_paragraph(document_or_cell, "", style=style)
    for run_data in runs:
        run = paragraph.add_run(run_data["text"])
        _apply_run_formatting(run, run_data)
    return paragraph


def _heading_style_for_type(item_type: str) -> str:
    key = item_type.strip().lower()
    if key == "epic":
        return STYLE_EPIC_HEADING
    if key == "feature":
        return STYLE_FEATURE_HEADING
    return STYLE_STORY_HEADING  # Story/User Story/Requirement/Issue/Test Case/... and anything deeper


def _load_base_document() -> Document:
    """Loads the org template and strips its sample body content, keeping
    only the trailing `sectPr` — that element carries the header/footer/page
    references, so everything subsequently added via `add_paragraph()` /
    `add_table()` (which python-docx always inserts before a trailing
    sectPr) renders with the template's logo, branding, and named styles
    intact. Falls back to a blank Document() if the template file is ever
    missing, so a packaging mistake never hard-crashes generation.
    """
    if not ORG_TEMPLATE_PATH.exists():
        logger.warning("org template not found at %s, falling back to a blank document", ORG_TEMPLATE_PATH)
        return Document()

    document = Document(str(ORG_TEMPLATE_PATH))
    body = document.element.body
    for child in list(body.iterchildren()):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return document


def _set_header_app_name(document: Document, app_name: str) -> None:
    """The header's "SRS For {App Name}" text lives inside a Content Control
    (`w:sdt`) — verified against the template's raw XML — which
    python-docx's `.paragraphs` doesn't traverse (it only walks direct
    `<w:p>` children, not ones nested inside `w:sdt`/text boxes). Iterating
    every `<w:t>` element in the header part directly finds it regardless of
    nesting. This is still safe, XML-escaped DOM manipulation — setting an
    lxml element's `.text` is exactly what python-docx's own `Run.text`
    setter does internally — not raw string templating.
    """
    if not app_name:
        return
    for section in document.sections:
        for t in section.header.part.element.iter(qn("w:t")):
            if t.text and _HEADER_APP_NAME_PLACEHOLDER in t.text:
                t.text = t.text.replace(_HEADER_APP_NAME_PLACEHOLDER, app_name)


def _set_footer_year(document: Document, year: str) -> None:
    """The footer's "Copyright © {year} ERA Info Tech Ltd." text has its
    year in its own isolated run (appears twice in the template's footer
    part) — same safe run-level `.text =` substitution as the header app
    name. Defaults to the current year (see `build_srs_document`) so a
    document generated in, say, 2027 doesn't ship a stale "2025" unless the
    user explicitly overrides it.
    """
    if not year:
        return
    for section in document.sections:
        for t in section.footer.part.element.iter(qn("w:t")):
            if t.text == _FOOTER_YEAR_PLACEHOLDER:
                t.text = year


def _set_footer_logo(document: Document, logo_bytes: bytes) -> None:
    """Swaps the footer's logo image in place. python-docx has no
    high-level "replace image" API — this overwrites the existing
    `ImagePart`'s blob directly (the standard workaround for this exact
    gap: `Part.blob` is read-only, so `_blob` is set directly, and
    `_image` is cleared so width/height/etc. get lazily re-derived from
    the new bytes instead of staying stale). Never raises — a branding
    cosmetic must never break document generation; on any failure this
    just leaves the template's baked-in ERA logo in place.
    """
    try:
        content_type = sniff_content_type(logo_bytes, "logo", None)
        for section in document.sections:
            image_rel = next((r for r in section.footer.part.rels.values() if r.reltype == RT.IMAGE), None)
            if image_rel is None:
                continue
            image_part = image_rel.target_part
            image_part._blob = logo_bytes  # noqa: SLF001 — see docstring, no public setter exists
            image_part._content_type = content_type  # noqa: SLF001
            image_part._image = None  # noqa: SLF001 — force lazy re-derivation from the new blob
    except Exception:  # noqa: BLE001
        logger.warning("could not swap footer logo, keeping the template default", exc_info=True)


def _add_toc_field(document: Document) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    r_element = run._r

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-4" \\h \\z \\u'
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose “Update Field” to populate this table of contents."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    r_element.append(fld_begin)
    r_element.append(instr)
    r_element.append(fld_separate)
    r_element.append(placeholder)
    r_element.append(fld_end)


def _add_two_col_table(document: Document, heading: str, rows: list[tuple[str, str]]) -> None:
    _safe_add_paragraph(document, heading, style=STYLE_SECTION_HEADING)
    table = document.add_table(rows=0, cols=2)
    table.style = TABLE_STYLE
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    document.add_paragraph()


def _add_history_table(document: Document, *, version: str, note: str, owner: str) -> None:
    _safe_add_paragraph(document, "Document History", style=STYLE_SECTION_HEADING)
    table = document.add_table(rows=1, cols=5)
    table.style = TABLE_STYLE
    header = table.rows[0].cells
    for idx, label in enumerate(["SL", "Date", "Version", "Description", "Created by"]):
        header[idx].text = label
    row = table.add_row().cells
    row[0].text = "01"
    row[1].text = date.today().isoformat()
    row[2].text = version
    row[3].text = note
    row[4].text = owner or "—"
    document.add_paragraph()


def _collect_features(roots: list[dict]) -> list[dict]:
    features: list[dict] = []

    def walk(node: dict) -> None:
        if node["work_item_type"].strip().lower() == "feature":
            features.append(node)
        for child in node.get("children", []):
            walk(child)

    for root in roots:
        walk(root)
    return features


def _add_feature_list_table(document: Document, roots: list[dict]) -> None:
    features = _collect_features(roots)
    if not features:
        return
    _safe_add_paragraph(document, "Feature List", style=STYLE_SECTION_HEADING)
    table = document.add_table(rows=1, cols=3)
    table.style = TABLE_STYLE
    header = table.rows[0].cells
    for idx, label in enumerate(["SL", "Feature", "Feature Details"]):
        header[idx].text = label
    for idx, feature in enumerate(features, start=1):
        row = table.add_row().cells
        row[0].text = str(idx)
        row[1].text = feature["title"]
        row[2].text = blocks_to_plain_text(feature.get("description_blocks") or [])[:500]
    document.add_paragraph()


def _add_cover_page(
    document: Document,
    *,
    app_name: str,
    document_metadata: dict,
    source_url: str,
    generated_at: str,
    snapshot_id: str,
) -> None:
    title1 = _safe_add_paragraph(document, "Software Requirement Specification (SRS)", style=STYLE_COVER_TITLE)
    title1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title2 = _safe_add_paragraph(document, app_name, style=STYLE_COVER_TITLE)
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER

    version = document_metadata.get("version") or "0.1"
    owner = document_metadata.get("owner") or ""
    status = document_metadata.get("status") or "Draft"

    try:
        generated_display = datetime.fromisoformat(generated_at).date().isoformat()
    except (ValueError, TypeError):
        generated_display = date.today().isoformat()

    for line in (
        f"Document Version {version}",
        f"Prepared by ERA Info Tech Ltd.{f' — {owner}' if owner else ''}",
        f"Date: {generated_display}",
    ):
        label = _safe_add_paragraph(document, line, style=STYLE_COVER_LABEL)
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph()

    _add_two_col_table(
        document,
        "Document Information",
        [
            ("Document ID", f"SRS-{snapshot_id[:8].upper()}"),
            ("Document Owner", owner or "—"),
            ("Date Submitted", date.today().isoformat()),
            ("Document Status", status),
            ("Document Version", version),
            ("Source", source_url),
        ],
    )
    _add_history_table(
        document,
        version=version,
        note=document_metadata.get("revision_note") or "Initial generation from Azure DevOps snapshot",
        owner=owner,
    )
    document.add_page_break()


# Every embed site displays images at a fixed ~110mm width regardless of
# source resolution — but the FULL-resolution bytes still get embedded in
# the file no matter how small it's displayed. Source screenshots/exported
# diagrams are routinely 2000-4000px wide, so a document with dozens of them
# balloons past what's safe: LibreOffice's PDF export can run out of memory
# rendering that many large images (reproduced live — a 30MB-of-images
# document aborted PDF export every time; downscaled to 12MB, it converted
# reliably), and Word has a well-known tendency to refuse or "repair"
# image-heavy documents past a certain size. 1600px is comfortably above
# what a 110mm embed ever needs even at high DPI/zoom, so this is a pure
# file-size fix with no visible quality loss.
_MAX_EMBEDDED_IMAGE_WIDTH_PX = 1600


def _downscale_image_bytes(image_bytes: bytes, *, max_width_px: int = _MAX_EMBEDDED_IMAGE_WIDTH_PX) -> bytes:
    """Every image is re-encoded, not just ones over `max_width_px` — a
    dense screenshot well under the width cap can still be a multi-MB PNG,
    and leaving it untouched was the actual reason a lot of images weren't
    getting any smaller. Format is never converted (PNG stays PNG, JPEG
    stays JPEG) — only the pixel/color encoding is tightened.
    """
    try:
        with Image.open(BytesIO(image_bytes)) as img:
            fmt = (img.format or "PNG").upper()  # .resize() below drops .format, so capture it first
            if img.width > max_width_px:
                ratio = max_width_px / img.width
                new_size = (max_width_px, max(1, round(img.height * ratio)))
                img = img.resize(new_size, Image.Resampling.LANCZOS)

            buf = BytesIO()
            if fmt == "PNG":
                # FASTOCTREE is the Pillow-recommended quantize method for
                # images with an alpha channel — a plain palette convert
                # silently flattens transparency to black, which would
                # corrupt any screenshot/diagram exported with a
                # transparent background.
                img.quantize(colors=256, method=Image.Quantize.FASTOCTREE).save(buf, format="PNG", optimize=True)
            elif fmt in ("JPEG", "JPG"):
                # Only force an RGB conversion when the source mode isn't
                # already JPEG-safe, so a grayscale JPEG doesn't get
                # needlessly blown up to 3-channel RGB.
                to_save = img if img.mode in ("RGB", "L") else img.convert("RGB")
                to_save.save(buf, format="JPEG", quality=80, optimize=True)
            else:
                img.save(buf, format=fmt, optimize=True)
            return buf.getvalue()
    except Exception:  # noqa: BLE001 — a bad/unreadable image falls back to the original bytes, not a crash
        logger.warning("failed to downscale image before embedding — using original bytes", exc_info=True)
        return image_bytes


def _insert_images(document: Document, minio, asset_refs: list[dict], *, width_mm: int = 110) -> None:
    for ref in asset_refs:
        try:
            image_bytes = _downscale_image_bytes(minio.download_bytes(ref["bucket"], ref["object_key"]))
            document.add_picture(BytesIO(image_bytes), width=Mm(width_mm))
        except Exception:  # noqa: BLE001 — one bad image must not break the whole render
            logger.warning("could not embed image %s into document", ref, exc_info=True)


_SOURCE_URL_KEY_LENGTH = 200  # matches asset_pipeline.DownloadJob's data-URI truncation


def _source_url_key(src: str | None) -> str:
    return (src or "")[:_SOURCE_URL_KEY_LENGTH]


def _render_image_block(
    document: Document, minio, block: dict, assets: list[dict], used_object_keys: set[str], *, caption: str | None
) -> None:
    """Renders one inline `<img>` placeholder from `html_to_blocks` by
    matching it back to the asset downloaded for it at import time — matched
    by `source_url` (truncated the same way on both sides), since that's the
    one stable identifier both the import-time download and this
    generation-time re-parse of the same HTML agree on, without needing any
    extra schema or bookkeeping in between.

    A caption (the heading immediately preceding this image in the source,
    e.g. "Context Diagram") renders directly above it — this is what makes
    each embedded diagram identifiable instead of an anonymous picture.
    """
    src_key = _source_url_key(block.get("src"))
    match = next(
        (
            asset
            for asset in assets
            if asset.get("object_key") not in used_object_keys and _source_url_key(asset.get("source_url")) == src_key
        ),
        None,
    )
    if match is None:
        # Not downloaded (or the match genuinely failed) — nothing to render
        # here. This asset, if it exists at all, is still eligible for the
        # end-of-node fallback pass, so it's never silently lost outright.
        return

    if caption:
        caption_para = document.add_paragraph()
        caption_run = caption_para.add_run(caption)
        caption_run.bold = True
        caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    try:
        image_bytes = _downscale_image_bytes(minio.download_bytes(match["bucket"], match["object_key"]))
        document.add_picture(BytesIO(image_bytes), width=Mm(110))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception:  # noqa: BLE001 — one bad image must not break the whole render
        logger.warning("could not embed image %s into document", match, exc_info=True)
    used_object_keys.add(match["object_key"])


def _render_table_block(document: Document, rows: list[list[list[dict]]]) -> None:
    """Real Azure DevOps custom fields (Data Dictionary especially) are
    genuinely authored as HTML tables — rendering them as a native docx
    table (not flattened prose) is what actually makes them readable; a
    flattened few-word-per-row table used to get stretched into unreadable,
    widely-spaced text by the org template's justified "S Notes" style.
    Each cell is itself a run list, so inline formatting inside a cell
    (rare but possible) still comes through.
    """
    if not rows:
        return
    num_cols = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=num_cols)
    table.style = TABLE_STYLE
    for row_idx, row in enumerate(rows):
        cells = table.add_row().cells
        for col_idx in range(num_cols):
            cell_runs = row[col_idx] if col_idx < len(row) else []
            paragraph = cells[col_idx].paragraphs[0]
            for run_data in cell_runs:
                run = paragraph.add_run(run_data["text"])
                _apply_run_formatting(run, run_data)
                if row_idx == 0:  # always bold the header row for readability
                    run.bold = True
    document.add_paragraph()


_CAPTION_MAX_LENGTH = 100  # a heading-like line (e.g. "Context Diagram"), not a paragraph of prose


# Sizes for headings that originate from Markdown-formatted custom field
# content (e.g. "## Overview" inside a "Data Dictionary" field) — bold,
# direct run formatting only, no named Word heading style. Reusing a real
# heading style (STYLE_SECTION_HEADING etc.) would pull every one of these
# into the auto-generated Table of Contents (_add_toc_field) alongside real
# Epic/Feature/Story headings, which isn't wanted for arbitrary org-authored
# field content — same reasoning as _render_content_sections's own
# direct-bold section label just below.
_MARKDOWN_HEADING_FONT_SIZE_PT = {1: 14, 2: 13, 3: 12}


def _render_heading_block(document: Document, runs: list[dict], level: int) -> None:
    paragraph = _add_runs_paragraph(document, runs, style=STYLE_NARRATIVE)
    size = _MARKDOWN_HEADING_FONT_SIZE_PT.get(level, 11)
    for run in paragraph.runs:
        run.bold = True
        run.font.size = Pt(size)


def _shade_paragraph(paragraph, fill_hex: str) -> None:
    """Paragraph background shading has no high-level python-docx API -
    this is the standard, safe way to do it: build a real `w:shd` element
    via lxml (same approach as everything else in this pipeline - never
    raw XML strings) and attach it to the paragraph's properties.
    """
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.append(shd)


def _render_code_block(document: Document, text: str, language: str | None) -> None:
    """A fenced ```code``` block from a Markdown custom field — grey-shaded,
    monospace, with exact line breaks preserved (a plain docx paragraph
    doesn't keep "\\n" as visual breaks on its own, hence the explicit
    add_break() per line rather than one run per block).
    """
    paragraph = document.add_paragraph()
    _shade_paragraph(paragraph, _CODE_BLOCK_SHADING_FILL)
    lines = text.split("\n") or [""]
    for i, line in enumerate(lines):
        run = paragraph.add_run(line or " ")  # a blank line still needs a run so the shading shows through it
        run.font.name = _CODE_FONT_NAME
        run.font.size = Pt(_CODE_BLOCK_FONT_SIZE_PT)
        if i < len(lines) - 1:
            run.add_break()  # soft line break - keeps every line inside the one shaded paragraph
    document.add_paragraph()  # breathing room after the block


def _style_exists(document: Document, name: str) -> bool:
    try:
        document.styles[name]
    except KeyError:
        return False
    return True


_LIST_INDENT_PER_LEVEL_PT = 18


def _list_style_for_level(document: Document, ordered: bool, level: int) -> str | None:
    """Word ships "List Bullet"/"List Number" for the top level and
    "... 2"/"... 3" for deeper nesting - prefer those native styles so a
    numbered outline with a bullet sub-list gets Word's own correct visual
    nesting. But the org template is a curated custom style set (see the
    STYLE_* constants above) that's only ever been verified to define "List
    Bullet" — NOT "List Number" or any "... 2"/"... 3" variant. Silently
    falling back to unstyled "Normal" paragraphs for those (as
    _safe_add_paragraph's generic KeyError handling would) would mean
    numbered/nested lists render with no bullet, no number, no indent at
    all — a real regression, not a graceful degradation. So this checks
    style existence up front and prefers, in order: the ideal per-level
    style, the base style with no level suffix, then any bullet at all
    (some visual list marker beats none) — real indentation is still
    applied directly per level regardless of which style wins, since that
    doesn't depend on the template defining anything extra.
    """
    base = STYLE_NUMBER if ordered else STYLE_BULLET
    candidates = ([f"{base} {level + 1}"] if level > 0 else []) + [base, STYLE_BULLET]
    for name in candidates:
        if _style_exists(document, name):
            return name
    return None


def _render_list_block(document: Document, block: dict, level: int = 0) -> None:
    ordered = bool(block.get("ordered"))
    style = _list_style_for_level(document, ordered, level)
    for item in block.get("items") or []:
        runs = item.get("runs") or []
        if runs:
            paragraph = _add_runs_paragraph(document, runs, style=style)
            if level > 0:
                paragraph.paragraph_format.left_indent = Pt(_LIST_INDENT_PER_LEVEL_PT * level)
        sublist = item.get("sublist")
        if sublist:
            _render_list_block(document, sublist, level=level + 1)


def _render_blocks(document: Document, blocks: list[dict], *, minio, assets: list[dict], used_object_keys: set[str]) -> None:
    """Renders `html_to_blocks` output with layout matching each block's
    actual shape: a real table for `table`, bullet paragraphs for `list`,
    an embedded picture for `image`, and left-aligned paragraphs for `text`
    — explicitly LEFT, overriding the org template's "S Notes" style
    default of JUSTIFY, which stretches short lines (a genuine table row
    once flattened to a handful of words, or any short line) into
    unreadable, widely-spaced text. Every run's bold/italic/underline
    formatting is preserved as-is.

    A short text block immediately followed by an image (e.g. "Context
    Diagram" right before its own `<img>`, matching how Azure's C4 tab is
    actually authored) renders as that image's caption instead of a
    separate paragraph — this is what makes each diagram identifiable and
    puts it where it actually belongs, instead of every image for a work
    item being dumped together at the end with no idea which was which.
    """
    index = 0
    total = len(blocks)
    while index < total:
        block = blocks[index]
        block_type = block.get("type")
        if block_type == "table":
            _render_table_block(document, block.get("rows") or [])
        elif block_type == "code":
            _render_code_block(document, block.get("text") or "", block.get("language"))
        elif block_type == "image":
            _render_image_block(document, minio, block, assets, used_object_keys, caption=None)
        elif block_type == "list":
            _render_list_block(document, block)
        elif block_type == "heading":
            runs = block.get("runs") or []
            if runs:
                _render_heading_block(document, runs, block.get("level") or 2)
        else:  # "text"
            runs = block.get("runs") or []
            if not runs:
                index += 1
                continue
            text_value = "".join(r["text"] for r in runs)
            next_block = blocks[index + 1] if index + 1 < total else None
            if next_block is not None and next_block.get("type") == "image" and len(text_value) <= _CAPTION_MAX_LENGTH:
                _render_image_block(document, minio, next_block, assets, used_object_keys, caption=text_value)
                index += 2
                continue
            # A manually-typed "- item" line (not a real <ul>) still reads
            # as a bullet, not a justified paragraph — strip just the
            # prefix off the first run, keeping its formatting intact.
            first_text = runs[0]["text"]
            stripped_first = first_text.lstrip("-*• \t")
            if stripped_first and stripped_first != first_text:
                runs = [{**runs[0], "text": stripped_first}, *runs[1:]]
                _add_runs_paragraph(document, runs, style=STYLE_BULLET)
            else:
                _add_runs_paragraph(document, runs, style=STYLE_NARRATIVE)
        index += 1


def _render_content_sections(
    document: Document, sections: list[dict], *, minio, assets: list[dict], used_object_keys: set[str]
) -> None:
    """Renders every custom-field content section (ERD, Class Diagram,
    Business Rules, Functional/Non-Functional Requirements, ...) as a
    labeled subsection. Discovered generically per work item — see
    srs_core.parsing.custom_fields — so this isn't hardcoded to any one
    organization's process template field names.
    """
    for section in sections:
        heading_para = document.add_paragraph()
        heading_run = heading_para.add_run(section["label"])
        heading_run.bold = True
        _render_blocks(
            document, section.get("blocks") or [], minio=minio, assets=assets, used_object_keys=used_object_keys
        )


def _render_node(document: Document, minio, node: dict, traceability: list[dict]) -> None:
    item_type = node["work_item_type"]
    assets = node.get("assets") or []
    # Tracks which downloaded assets got matched to an in-place <img>
    # placeholder — an embedded diagram positioned right after its own
    # heading. Formal "AttachedFile" attachments never appear inline in any
    # field's HTML at all, so they never match anything here and correctly
    # fall through to the batch-insert below, same as before this feature.
    used_object_keys: set[str] = set()

    if item_type.lower() in _LEAF_TYPES:
        state_note = f" ({node['state']})" if node.get("state") else ""
        _safe_add_paragraph(
            document, f"{item_type} #{node['azure_work_item_id']}: {node['title']}{state_note}", style=STYLE_BULLET
        )
        # A Task/Bug renders as a single bullet line, not a full section, but
        # any image attached to it still has to be captured — a bug once had
        # a real diagram fetched and stored correctly, then silently dropped
        # here because only heading-level nodes inserted images.
        if node.get("content_sections"):
            _render_content_sections(
                document, node["content_sections"], minio=minio, assets=assets, used_object_keys=used_object_keys
            )
        leftover = [a for a in assets if a.get("object_key") not in used_object_keys]
        if leftover:
            _insert_images(document, minio, leftover, width_mm=90)
        traceability.append(
            {
                "type": item_type,
                "azure_id": node["azure_work_item_id"],
                "title": node["title"],
                "azure_url": node["azure_url"],
            }
        )
        return

    _safe_add_paragraph(
        document,
        f"{item_type} #{node['azure_work_item_id']}: {node['title']}",
        style=_heading_style_for_type(item_type),
    )

    meta_para = document.add_paragraph()
    meta_run = meta_para.add_run(f"Status: {node['state']}")
    meta_run.italic = True
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    if node.get("description_blocks"):
        _render_blocks(
            document, node["description_blocks"], minio=minio, assets=assets, used_object_keys=used_object_keys
        )

    if node.get("acceptance_criteria_blocks"):
        ac_heading = document.add_paragraph()
        ac_run = ac_heading.add_run("Acceptance Criteria")
        ac_run.bold = True
        _render_blocks(
            document, node["acceptance_criteria_blocks"], minio=minio, assets=assets, used_object_keys=used_object_keys
        )

    if node.get("content_sections"):
        _render_content_sections(
            document, node["content_sections"], minio=minio, assets=assets, used_object_keys=used_object_keys
        )

    # Anything never matched to an in-place placeholder — a formal
    # attachment, or an embedded image html_to_blocks couldn't re-locate —
    # still renders here, so nothing is ever silently dropped.
    leftover = [a for a in assets if a.get("object_key") not in used_object_keys]
    if leftover:
        _insert_images(document, minio, leftover)

    traceability.append(
        {
            "type": item_type,
            "azure_id": node["azure_work_item_id"],
            "title": node["title"],
            "azure_url": node["azure_url"],
        }
    )

    for child in node.get("children", []):
        _render_node(document, minio, child, traceability)


def _add_traceability_table(document: Document, traceability: list[dict]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.style = TABLE_STYLE
    header = table.rows[0].cells
    header[0].text = "ID"
    header[1].text = "Type"
    header[2].text = "Title"
    header[3].text = "Source"
    for entry in traceability:
        row = table.add_row().cells
        row[0].text = f"#{entry['azure_id']}"
        row[1].text = entry["type"]
        row[2].text = entry["title"]
        row[3].text = entry["azure_url"]


def resolve_document_title(context: dict) -> str:
    document_metadata = context.get("document_metadata") or {}
    roots = context.get("roots") or []
    epic_titles = [r["title"] for r in roots if r["work_item_type"] == "Epic"]
    return document_metadata.get("app_name") or (epic_titles[0] if epic_titles else "SRS Document")


def build_srs_document(context: dict, minio) -> bytes:
    document = _load_base_document()

    document_metadata = context.get("document_metadata") or {}
    roots = context["roots"]
    epic_titles = [r["title"] for r in roots if r["work_item_type"] == "Epic"]
    app_name = resolve_document_title(context)

    _set_header_app_name(document, app_name)
    _set_footer_year(document, document_metadata.get("footer_year") or str(date.today().year))
    logo_override = context.get("logo_override")
    if logo_override:
        try:
            logo_bytes = minio.download_bytes(logo_override["bucket"], logo_override["object_key"])
            _set_footer_logo(document, logo_bytes)
        except Exception:  # noqa: BLE001 — a branding cosmetic must never break generation
            logger.warning("could not download logo override, keeping the template default", exc_info=True)
    _add_cover_page(
        document,
        app_name=app_name,
        document_metadata=document_metadata,
        source_url=context["source_url"],
        generated_at=context["generated_at"],
        snapshot_id=context["snapshot_id"],
    )

    _safe_add_paragraph(document, "Table of Contents", style=STYLE_SECTION_HEADING)
    _add_toc_field(document)
    document.add_page_break()

    _safe_add_paragraph(document, "Introduction", style=STYLE_SECTION_HEADING)
    if context.get("ai_introduction"):
        note = document.add_paragraph()
        note_run = note.add_run("AI-generated summary, based on the imported backlog — verify against source data.")
        note_run.italic = True
        note_run.font.size = Pt(8)
        note_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        _safe_add_paragraph(document, context["ai_introduction"], style=STYLE_NARRATIVE)
    else:
        _safe_add_paragraph(
            document,
            f"This document describes the requirements captured from the Azure DevOps backlog at "
            f"{context['source_url']}, imported and sealed as snapshot {context['snapshot_id']}.",
            style=STYLE_NARRATIVE,
        )

    _safe_add_paragraph(document, "Scope", style=STYLE_SECTION_HEADING)
    if epic_titles:
        _safe_add_paragraph(document, "This specification covers the following epics:", style=STYLE_NARRATIVE)
        for title in epic_titles:
            _safe_add_paragraph(document, title, style=STYLE_BULLET)
    else:
        _safe_add_paragraph(
            document, "This specification covers the imported and selected backlog items below.", style=STYLE_NARRATIVE
        )

    _add_feature_list_table(document, roots)

    _safe_add_paragraph(document, "Requirements", style=STYLE_SECTION_HEADING)
    traceability: list[dict] = []
    for root in roots:
        _render_node(document, minio, root, traceability)

    _safe_add_paragraph(document, "Traceability Matrix", style=STYLE_SECTION_HEADING)
    _add_traceability_table(document, traceability)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()

"""One-off generator for templates/default_srs_template.docx.

Not part of the runtime — run once (or whenever the template's Jinja tags
need to change) to (re)produce the committed binary template. Phase 3
replaces this hardcoded template with org-managed, uploaded, versioned ones;
Phase 1 renders against exactly this one file.
"""

import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def build(output_path: str) -> None:
    doc = Document()

    title = doc.add_heading("Software Requirements Specification", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("Source: {{ source_url }}")
    doc.add_paragraph("Generated: {{ generated_at }}")
    doc.add_paragraph("Snapshot: {{ snapshot_id }}")

    doc.add_heading("Work Items", level=1)
    doc.add_paragraph("{% for item in work_items %}")

    p = doc.add_paragraph()
    run = p.add_run("{{ item.work_item_type }} #{{ item.azure_work_item_id }}: {{ item.title }} ({{ item.state }})")
    run.bold = True
    run.font.size = Pt(11)

    doc.add_paragraph("{% if item.description %}{{ item.description }}{% endif %}")
    doc.add_paragraph("{% if item.image %}{{ item.image }}{% endif %}")
    doc.add_paragraph("")
    doc.add_paragraph("{% endfor %}")

    doc.add_heading("Traceability", level=1)
    doc.add_paragraph("{% for item in work_items %}")
    doc.add_paragraph("{{ item.azure_work_item_id }} -> {{ item.azure_url }}")
    doc.add_paragraph("{% endfor %}")

    doc.save(output_path)


if __name__ == "__main__":
    build(sys.argv[1] if len(sys.argv) > 1 else "default_srs_template.docx")

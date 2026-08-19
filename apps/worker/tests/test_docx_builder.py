"""Unit tests for the python-docx-based render engine, independent of the
full Celery pipeline — build_srs_document is a pure function of
(context, minio_client).
"""

import struct
import zipfile
import zlib
from io import BytesIO
from unittest.mock import MagicMock

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

import src.tasks.docx_builder as docx_builder_module
from src.tasks.docx_builder import build_srs_document


def _media_file_count(docx_bytes: bytes) -> int:
    zf = zipfile.ZipFile(BytesIO(docx_bytes))
    return len([n for n in zf.namelist() if n.startswith("word/media/")])


def _header_text(doc: Document) -> str:
    # The org template's header text lives inside a Content Control
    # (w:sdt), which python-docx's high-level `.paragraphs` doesn't
    # traverse — walking every <w:t> in the part directly finds it
    # regardless of nesting, same technique docx_builder itself uses.
    return "".join(t.text or "" for t in doc.sections[0].header.part.element.iter(qn("w:t")))


def _footer_text(doc: Document) -> str:
    # Same story for the footer: its text lives inside a floating text box.
    return "".join(t.text or "" for t in doc.sections[0].footer.part.element.iter(qn("w:t")))


def _png_chunk(chunk_type: bytes, data: bytes) -> bytes:
    return struct.pack(">I", len(data)) + chunk_type + data + struct.pack(">I", zlib.crc32(chunk_type + data))


def _build_minimal_png(pixel_value: int = 0) -> bytes:
    # add_picture parses real PNG chunk structure — a bare signature isn't enough.
    signature = b"\x89PNG\r\n\x1a\n"
    ihdr = _png_chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 0, 0, 0, 0))
    idat = _png_chunk(b"IDAT", zlib.compress(bytes([0, pixel_value])))
    iend = _png_chunk(b"IEND", b"")
    return signature + ihdr + idat + iend


FAKE_PNG = _build_minimal_png()


def _fake_minio(image_bytes: bytes = FAKE_PNG) -> MagicMock:
    minio = MagicMock()
    minio.download_bytes.return_value = image_bytes
    return minio


def _base_context(roots: list[dict]) -> dict:
    return {
        "source_url": "https://dev.azure.com/org/proj",
        "generated_at": "2026-01-01T00:00:00Z",
        "snapshot_id": "snap-1",
        "total_count": len(roots),
        "ai_introduction": None,
        "roots": roots,
    }


def _node(work_item_type, azure_id, title, **overrides):
    node = {
        "work_item_type": work_item_type,
        "azure_work_item_id": azure_id,
        "title": title,
        "state": "New",
        "description_blocks": [],
        "acceptance_criteria_blocks": [],
        "content_sections": [],
        "azure_url": f"https://dev.azure.com/org/proj/_workitems/edit/{azure_id}",
        "assets": [],
        "children": [],
    }
    node.update(overrides)
    return node


def _run(text: str, *, bold: bool = False, italic: bool = False, underline: bool = False) -> dict:
    return {"text": text, "bold": bold, "italic": italic, "underline": underline}


def _text_blocks(text: str, **run_kwargs) -> list[dict]:
    return [{"type": "text", "runs": [_run(text, **run_kwargs)]}]


def _content_section(label: str, text: str) -> dict:
    return {"label": label, "blocks": _text_blocks(text)}


def _table_block(rows: list[list[str]]) -> dict:
    return {"type": "table", "rows": [[[_run(cell)] for cell in row] for row in rows]}


def _list_block(items: list[str]) -> dict:
    return {"type": "list", "items": [[_run(item)] for item in items]}


def _asset(bucket: str, object_key: str, source_url: str | None = None) -> dict:
    return {"bucket": bucket, "object_key": object_key, "source_url": source_url}


def _image_block(src: str) -> dict:
    return {"type": "image", "src": src}


def _captioned_image(caption: str, src: str) -> list[dict]:
    return [{"type": "text", "runs": [_run(caption)]}, _image_block(src)]


def test_task_leaf_node_still_embeds_its_own_images():
    """Regression test for a real bug: a Task/Bug renders as a single
    bullet line (not a full section), and images attached to it were
    silently dropped because only heading-level nodes inserted images —
    even though the image had been correctly fetched and stored.
    """
    task = _node("Task", 2, "UI/UX Design", assets=[{"bucket": "b", "object_key": "k"}])
    epic = _node("Epic", 1, "Parent Epic", children=[task])

    # Baseline: the org template itself carries its own media (the footer
    # logo, plus an orphaned image left over from the sample body content
    # that gets stripped) — count against that baseline, not an assumed 0.
    baseline = _media_file_count(build_srs_document(_base_context([_node("Epic", 9, "No assets")]), _fake_minio()))

    minio = _fake_minio()
    docx_bytes = build_srs_document(_base_context([epic]), minio)

    assert _media_file_count(docx_bytes) == baseline + 1
    minio.download_bytes.assert_called_once_with("b", "k")


def test_multiple_images_on_one_node_all_embedded():
    # Distinct byte content per image — python-docx correctly deduplicates
    # byte-identical images into one shared media file, so using the same
    # fake bytes three times would legitimately produce only one file.
    distinct_pngs = [_build_minimal_png(pixel_value=i + 1) for i in range(3)]
    epic = _node(
        "Epic",
        1,
        "Epic with many diagrams",
        assets=[{"bucket": "b", "object_key": f"k{i}"} for i in range(3)],
    )

    baseline = _media_file_count(build_srs_document(_base_context([_node("Epic", 9, "No assets")]), _fake_minio()))

    minio = MagicMock()
    minio.download_bytes.side_effect = distinct_pngs
    docx_bytes = build_srs_document(_base_context([epic]), minio)

    assert _media_file_count(docx_bytes) == baseline + 3


def test_hierarchy_headings_use_the_org_templates_named_styles():
    """Epic/Feature/Story map to the org template's real style names (not
    generic Word "Heading N" levels) — verified against the actual template
    file: "S Heading" is used at Epic level, "S Heading 2" for Features
    (matching e.g. "Feature-001: Authentication" in the source template),
    "S Heading 3" for Stories (matching "User Story: ...").
    """
    story = _node("User Story", 3, "As a user...")
    feature = _node("Feature", 2, "Signup flow", children=[story])
    epic = _node("Epic", 1, "Onboarding", children=[feature])

    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    # Restricted to "S Heading*" styles specifically — the Scope section
    # also lists each Epic's bare title as a "List Bullet" bullet, which
    # would otherwise be a false match for a substring search on "Onboarding".
    styles_by_text = {
        p.text: p.style.name for p in doc.paragraphs if p.style and p.style.name.startswith("S Heading")
    }
    epic_heading = next(text for text in styles_by_text if "Onboarding" in text)
    feature_heading = next(text for text in styles_by_text if "Signup flow" in text)
    story_heading = next(text for text in styles_by_text if "As a user" in text)

    assert styles_by_text[epic_heading] == "S Heading"
    assert styles_by_text[feature_heading] == "S Heading 2"
    assert styles_by_text[story_heading] == "S Heading 3"


def test_task_leaf_uses_list_bullet_style_with_real_bullet_glyphs():
    """Regression test for a real reported bug: bullets rendered as plain,
    unmarked paragraphs — no bullet glyph at all. Root cause was
    STYLE_BULLET pointing at "List Paragraph", which (verified against the
    org template's styles.xml) carries no <w:numPr> anywhere in its
    definition; "List Bullet" is the style Word defines WITH numbering
    baked in. Checking the style name alone isn't a strong enough
    assertion — this also confirms the resolved style itself actually
    carries a numbering reference, so a future style-name swap can't
    silently reintroduce the same bug.
    """
    task = _node("Task", 2, "Some task")
    epic = _node("Epic", 1, "Epic", children=[task])
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    task_para = next(p for p in doc.paragraphs if "Some task" in p.text)
    assert task_para.style.name == "List Bullet"
    numpr = task_para.style.element.find(".//" + qn("w:numPr"))
    assert numpr is not None


def test_org_template_logo_and_branding_carry_over_to_every_generated_document():
    """The whole point of using the org template as the base document: its
    footer (with the ERA logo image) and "CONFIDENTIAL" header must survive
    into every generated document, not just the sample BankAsia one.
    """
    epic = _node("Epic", 1, "Any Project")
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    assert "ERA Info Tech" in _footer_text(doc)
    assert "CONFIDENTIAL" in _header_text(doc)

    zf = zipfile.ZipFile(BytesIO(docx_bytes))
    assert any(name.startswith("word/media/") for name in zf.namelist())


def _footer_image_blob(doc: Document) -> bytes:
    image_rel = next(r for r in doc.sections[0].footer.part.rels.values() if r.reltype == RT.IMAGE)
    return image_rel.target_part.blob


def test_logo_override_replaces_the_footer_image():
    """Org branding lets a tenant replace the document logo from the UI —
    the override must actually land in the footer's image part, not just
    get downloaded and discarded.
    """
    context = _base_context([_node("Epic", 1, "E")])
    context["logo_override"] = {"bucket": "b", "object_key": "custom-logo"}
    minio = _fake_minio()  # returns FAKE_PNG for any download_bytes call
    docx_bytes = build_srs_document(context, minio)
    doc = Document(BytesIO(docx_bytes))

    assert _footer_image_blob(doc) == FAKE_PNG
    minio.download_bytes.assert_any_call("b", "custom-logo")


def test_no_logo_override_keeps_the_templates_default_logo():
    docx_bytes = build_srs_document(_base_context([_node("Epic", 1, "E")]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    blob = _footer_image_blob(doc)
    assert blob != FAKE_PNG  # still the template's own ERA logo bytes
    assert len(blob) > 0


def test_logo_override_download_failure_falls_back_without_crashing():
    context = _base_context([_node("Epic", 1, "E")])
    context["logo_override"] = {"bucket": "b", "object_key": "missing"}
    minio = MagicMock()
    minio.download_bytes.side_effect = RuntimeError("network error")

    docx_bytes = build_srs_document(context, minio)  # must not raise
    doc = Document(BytesIO(docx_bytes))
    assert "ERA Info Tech" in _footer_text(doc)
    assert _footer_image_blob(doc) != FAKE_PNG  # untouched template default


def test_header_app_name_is_substituted_from_document_metadata():
    context = _base_context([_node("Epic", 1, "E")])
    context["document_metadata"] = {"app_name": "Loan Approval and Management System"}
    docx_bytes = build_srs_document(context, _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    header_text = _header_text(doc)
    assert "Loan Approval and Management System" in header_text
    assert "Mobile Banking App" not in header_text


def test_footer_year_is_substituted_from_document_metadata():
    context = _base_context([_node("Epic", 1, "E")])
    context["document_metadata"] = {"footer_year": "2030"}
    docx_bytes = build_srs_document(context, _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    footer_text = _footer_text(doc)
    assert "2030" in footer_text
    assert "2025" not in footer_text


def test_footer_year_defaults_to_the_current_year_when_not_overridden():
    from datetime import date

    docx_bytes = build_srs_document(_base_context([_node("Epic", 1, "E")]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    footer_text = _footer_text(doc)
    assert str(date.today().year) in footer_text
    assert "2025" not in footer_text or str(date.today().year) == "2025"


def test_document_metadata_populates_front_matter_and_falls_back_when_absent():
    epic = _node("Epic", 1, "Fallback App")
    context = _base_context([epic])
    context["document_metadata"] = {
        "app_name": "Real App Name",
        "version": "1.2",
        "owner": "Jane Doe",
        "status": "Final",
        "revision_note": "Second pass",
    }
    docx_bytes = build_srs_document(context, _fake_minio())
    full_text = "\n".join(p.text for p in Document(BytesIO(docx_bytes)).paragraphs)
    tables_text = "\n".join(cell.text for t in Document(BytesIO(docx_bytes)).tables for row in t.rows for cell in row.cells)

    assert "Real App Name" in full_text
    assert "Jane Doe" in tables_text
    assert "Final" in tables_text
    assert "1.2" in tables_text
    assert "Second pass" in tables_text

    # No document_metadata at all must never crash generation — falls back
    # to the first Epic's title as the app name.
    fallback_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    fallback_text = "\n".join(p.text for p in Document(BytesIO(fallback_bytes)).paragraphs)
    assert "Fallback App" in fallback_text


def test_feature_list_table_populated_from_selected_features():
    story = _node("User Story", 3, "A story")
    feature = _node(
        "Feature", 2, "Signup flow", description_blocks=_text_blocks("Lets a user create an account"), children=[story]
    )
    epic = _node("Epic", 1, "Onboarding", children=[feature])

    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    tables_text = [
        "\n".join(cell.text for row in t.rows for cell in row.cells) for t in doc.tables
    ]
    assert any("Signup flow" in t and "Lets a user create an account" in t for t in tables_text)


def test_acceptance_criteria_rendered_as_bullets():
    # Two separate text blocks — matching how html_to_blocks actually
    # splits two dash-prefixed lines (each its own <div>/<p> in real HTML)
    # into two independent blocks, not one block with an embedded newline.
    story = _node(
        "User Story", 1, "Story", acceptance_criteria_blocks=_text_blocks("- must do X") + _text_blocks("- must do Y")
    )
    docx_bytes = build_srs_document(_base_context([story]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "must do X" in full_text
    assert "must do Y" in full_text

    bullet_paras = [p.text for p in doc.paragraphs if p.style and p.style.name == "List Bullet"]
    assert "must do X" in bullet_paras
    assert "must do Y" in bullet_paras


def test_acceptance_criteria_real_list_markup_rendered_as_bullets():
    """A genuine <ul><li> (parsed upstream into a "list" block), not just a
    manually dash-prefixed line, must also render as real bullets.
    """
    story = _node(
        "User Story",
        1,
        "Story",
        acceptance_criteria_blocks=[_list_block(["must do X", "must do Y"])],
    )
    docx_bytes = build_srs_document(_base_context([story]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))
    bullet_paras = [p.text for p in doc.paragraphs if p.style and p.style.name == "List Bullet"]
    assert "must do X" in bullet_paras
    assert "must do Y" in bullet_paras


def test_traceability_table_includes_every_node_including_tasks():
    task = _node("Task", 2, "Some task")
    epic = _node("Epic", 1, "Epic", children=[task])

    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    table = doc.tables[-1]
    ids_in_table = {row.cells[0].text for row in table.rows[1:]}
    assert "#1" in ids_in_table
    assert "#2" in ids_in_table


def _build_oversized_png(width: int, height: int) -> bytes:
    from PIL import Image

    img = Image.new("RGB", (width, height), color=(120, 60, 200))
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def _build_test_jpeg(width: int, height: int) -> bytes:
    from PIL import Image

    img = Image.new("RGB", (width, height), color=(200, 90, 40))
    buf = BytesIO()
    img.save(buf, format="JPEG")
    return buf.getvalue()


def _build_transparent_png(width: int, height: int) -> bytes:
    from PIL import Image

    img = Image.new("RGBA", (width, height), color=(120, 60, 200, 255))
    # Punch out a fully transparent region so quantization has real alpha to preserve.
    for x in range(width // 2):
        for y in range(height // 2):
            img.putpixel((x, y), (0, 0, 0, 0))
    buf = BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def test_downscale_image_bytes_shrinks_images_over_the_max_width():
    from PIL import Image

    from src.tasks.docx_builder import _MAX_EMBEDDED_IMAGE_WIDTH_PX, _downscale_image_bytes

    oversized = _build_oversized_png(3000, 1500)
    result = _downscale_image_bytes(oversized)

    with Image.open(BytesIO(result)) as img:
        assert img.width == _MAX_EMBEDDED_IMAGE_WIDTH_PX
        assert img.height == 800  # 2:1 aspect ratio preserved (3000x1500 -> 1600x800)
        assert img.format == "PNG"
    assert len(result) < len(oversized)


def test_downscale_image_bytes_reencodes_small_png_instead_of_passing_it_through():
    from PIL import Image

    from src.tasks.docx_builder import _downscale_image_bytes

    small = _build_oversized_png(400, 300)
    result = _downscale_image_bytes(small)

    assert result != small  # every image is re-encoded now, not just oversized ones
    with Image.open(BytesIO(result)) as img:
        assert (img.width, img.height) == (400, 300)  # dimensions unchanged
        assert img.format == "PNG"  # format unchanged


def test_downscale_image_bytes_preserves_jpeg_format_when_small():
    from PIL import Image

    from src.tasks.docx_builder import _downscale_image_bytes

    small = _build_test_jpeg(400, 300)
    result = _downscale_image_bytes(small)

    with Image.open(BytesIO(result)) as img:
        assert img.format == "JPEG"  # never converted to PNG
        assert (img.width, img.height) == (400, 300)


def test_downscale_image_bytes_preserves_jpeg_format_and_aspect_ratio_when_large():
    from PIL import Image

    from src.tasks.docx_builder import _MAX_EMBEDDED_IMAGE_WIDTH_PX, _downscale_image_bytes

    oversized = _build_test_jpeg(3000, 1500)
    result = _downscale_image_bytes(oversized)

    with Image.open(BytesIO(result)) as img:
        assert img.format == "JPEG"
        assert img.width == _MAX_EMBEDDED_IMAGE_WIDTH_PX
        assert img.height == 800  # 2:1 aspect ratio preserved


def test_downscale_image_bytes_preserves_transparency_through_png_quantization():
    from PIL import Image

    from src.tasks.docx_builder import _downscale_image_bytes

    transparent = _build_transparent_png(200, 200)
    result = _downscale_image_bytes(transparent)

    with Image.open(BytesIO(result)) as img:
        assert img.format == "PNG"
        has_alpha_channel = img.mode in ("RGBA", "LA") or "transparency" in img.info
        assert has_alpha_channel
        rgba = img.convert("RGBA")
        assert rgba.getpixel((0, 0))[3] == 0  # the punched-out corner is still transparent
        assert rgba.getpixel((150, 150))[3] > 0  # the opaque region is still opaque


def test_downscale_image_bytes_falls_back_to_original_on_unreadable_input():
    from src.tasks.docx_builder import _downscale_image_bytes

    garbage = b"not an image at all"
    assert _downscale_image_bytes(garbage) == garbage


def test_large_embedded_image_is_downscaled_in_the_actual_render_path():
    """Reproduces the real failure live-verified this session: a document
    with several multi-megabyte, wide screenshots aborted LibreOffice's PDF
    export every time; downscaled, it converted reliably. This checks the
    render path actually applies that downscale, not just the helper in
    isolation.
    """
    oversized = _build_oversized_png(3000, 1500)
    baseline_names = set(
        zipfile.ZipFile(
            BytesIO(build_srs_document(_base_context([_node("Epic", 9, "No assets")]), _fake_minio()))
        ).namelist()
    )

    epic = _node("Epic", 1, "Epic", assets=[{"bucket": "b", "object_key": "big.png"}])
    minio = _fake_minio(image_bytes=oversized)
    docx_bytes = build_srs_document(_base_context([epic]), minio)

    zf = zipfile.ZipFile(BytesIO(docx_bytes))
    new_media = [n for n in zf.namelist() if n.startswith("word/media/") and n not in baseline_names]
    assert new_media, "expected the epic's image to be embedded as a new media file"
    from PIL import Image

    with Image.open(BytesIO(zf.read(new_media[0]))) as img:
        assert img.width <= 1600


def test_a_broken_image_is_skipped_without_failing_the_whole_render():
    epic = _node("Epic", 1, "Epic", assets=[{"bucket": "b", "object_key": "bad"}])
    minio = MagicMock()
    minio.download_bytes.side_effect = RuntimeError("network error")

    docx_bytes = build_srs_document(_base_context([epic]), minio)  # must not raise
    doc = Document(BytesIO(docx_bytes))
    assert "Epic" in "\n".join(p.text for p in doc.paragraphs)


def test_custom_content_sections_rendered_on_heading_nodes():
    """Custom process-template fields (ERD, Class Diagram, Business Rules,
    Functional/Non-Functional Requirements, ...) hold real requirement
    content this org's process captures outside the standard
    Description/AcceptanceCriteria fields — they must actually show up in
    the document, not just get fetched and silently dropped.
    """
    epic = _node(
        "Epic",
        1,
        "AI Powered Assessment Agent",
        content_sections=[
            _content_section("Business Rules", "Only authenticated users may access the agent."),
            _content_section("ERD", "See the entity relationship diagram below."),
        ],
    )
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    full_text = "\n".join(p.text for p in Document(BytesIO(docx_bytes)).paragraphs)

    assert "Business Rules" in full_text
    assert "Only authenticated users may access the agent." in full_text
    assert "ERD" in full_text
    assert "See the entity relationship diagram below." in full_text


def test_custom_content_sections_rendered_on_leaf_task_nodes_too():
    task = _node(
        "Task",
        2,
        "Some task",
        content_sections=[_content_section("Notes", "Task-level detail that must not be dropped.")],
    )
    epic = _node("Epic", 1, "Parent", children=[task])
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    full_text = "\n".join(p.text for p in Document(BytesIO(docx_bytes)).paragraphs)

    assert "Task-level detail that must not be dropped." in full_text


def test_ai_introduction_used_when_present_with_visible_marker():
    context = _base_context([_node("Epic", 1, "E")])
    context["ai_introduction"] = "Custom AI summary text."
    docx_bytes = build_srs_document(context, _fake_minio())
    full_text = "\n".join(p.text for p in Document(BytesIO(docx_bytes)).paragraphs)
    assert "Custom AI summary text." in full_text
    assert "AI-generated summary" in full_text


def test_falls_back_to_a_blank_document_if_the_org_template_is_missing(monkeypatch, tmp_path):
    """A packaging mistake that drops the template file must never hard-crash
    generation — falls back to a blank Document() instead.
    """
    monkeypatch.setattr(docx_builder_module, "ORG_TEMPLATE_PATH", tmp_path / "does-not-exist.docx")
    epic = _node("Epic", 1, "Still Works")
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())  # must not raise
    full_text = "\n".join(p.text for p in Document(BytesIO(docx_bytes)).paragraphs)
    assert "Still Works" in full_text


def test_table_content_section_renders_as_a_real_docx_table_not_flattened_text():
    """Regression test for a real reported bug: custom fields like "Data
    Dictionary" are genuinely authored as HTML tables in Azure DevOps.
    Flattening a table to plain text turned each row into a sparse
    few-word paragraph that the org template's justified "S Notes" style
    then stretched into unreadable, widely spaced text. Table blocks must
    render as an actual docx table instead.
    """
    epic = _node(
        "Epic",
        1,
        "Epic",
        content_sections=[
            {
                "label": "Data Dictionary",
                "blocks": [
                    _table_block(
                        [
                            ["Field Name", "Type", "Description"],
                            ["customer_id", "UUID", "Primary key for the customer record"],
                            ["kyc_status", "ENUM", "Current KYC verification state"],
                        ]
                    )
                ],
            }
        ],
    )
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    tables_text = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    assert any(
        rows[0] == ["Field Name", "Type", "Description"]
        and ["customer_id", "UUID", "Primary key for the customer record"] in rows
        for rows in tables_text
    )

    # It must NOT also appear as flattened prose paragraphs.
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Field NameTypeDescription" not in full_text.replace(" ", "")


def test_table_header_row_is_bolded():
    epic = _node(
        "Epic",
        1,
        "Epic",
        content_sections=[{"label": "Data Dictionary", "blocks": [_table_block([["Field", "Type"], ["id", "UUID"]])]}],
    )
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    table = next(t for t in doc.tables if t.rows[0].cells[0].text == "Field")
    header_run = table.rows[0].cells[0].paragraphs[0].runs[0]
    data_run = table.rows[1].cells[0].paragraphs[0].runs[0]
    assert header_run.bold is True
    assert not data_run.bold


def test_list_content_section_renders_each_item_as_a_separate_bullet():
    epic = _node(
        "Epic",
        1,
        "Epic",
        content_sections=[
            {"label": "Business Rules", "blocks": [_list_block(["Only admins may approve.", "Every change is logged."])]}
        ],
    )
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    bullet_paras = [p.text for p in doc.paragraphs if p.style and p.style.name == "List Bullet"]
    assert "Only admins may approve." in bullet_paras
    assert "Every change is logged." in bullet_paras


def test_narrative_content_is_left_aligned_not_justified():
    """Regression test for a real reported bug: the org template's "S
    Notes" style defaults to JUSTIFY, which stretched short lines (a table
    row flattened to a few words, or any short line) into unreadable,
    widely spaced text. Every generated paragraph must be explicitly LEFT
    aligned regardless of the underlying style's default.
    """
    epic = _node("Epic", 1, "Epic", description_blocks=_text_blocks("A short line."))
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    short_para = next(p for p in doc.paragraphs if p.text == "A short line.")
    assert short_para.style.name == "S Notes"  # still uses the narrative style...
    assert short_para.alignment == WD_ALIGN_PARAGRAPH.LEFT  # ...but never justified


def test_inline_bold_and_italic_runs_are_preserved_in_the_rendered_paragraph():
    """Regression test for a real reported bug: a work item's inline bold
    text (e.g. "using either the **Short CIF** or **Full CIF**") rendered
    as plain, unstyled text — the run-level bold/italic/underline flags
    from html_to_blocks must actually reach the docx run objects.
    """
    epic = _node(
        "Epic",
        1,
        "Epic",
        description_blocks=[
            {
                "type": "text",
                "runs": [
                    _run("Enable login using either the "),
                    _run("Short CIF", bold=True),
                    _run(" or "),
                    _run("Full CIF", bold=True),
                    _run(" onboarding process."),
                ],
            }
        ],
    )
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    target = next(p for p in doc.paragraphs if "Enable login using either the" in p.text)
    runs_by_text = {r.text: r.bold for r in target.runs}
    assert runs_by_text["Short CIF"] is True
    assert runs_by_text["Full CIF"] is True
    assert runs_by_text["Enable login using either the "] is not True


def test_inline_bold_preserved_in_bullet_lists_too():
    epic = _node(
        "Epic",
        1,
        "Epic",
        content_sections=[
            {
                "label": "Business Rules",
                "blocks": [{"type": "list", "items": [[_run("Only "), _run("admins", bold=True), _run(" may approve.")]]}],
            }
        ],
    )
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    target = next(p for p in doc.paragraphs if "Only " in p.text and "may approve" in p.text)
    runs_by_text = {r.text: r.bold for r in target.runs}
    assert runs_by_text["admins"] is True


def test_description_blocks_preserve_table_structure():
    """Not just custom fields — a work item's own Description field can
    also be an HTML table (e.g. a Requirement Analysis item authored as a
    comparison table), and must render the same way.
    """
    epic = _node(
        "Epic",
        1,
        "Requirement Analysis",
        description_blocks=[_table_block([["Requirement", "Priority"], ["Login via SSO", "High"]])],
    )
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())
    doc = Document(BytesIO(docx_bytes))

    tables_text = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    assert any(["Login via SSO", "High"] in rows for rows in tables_text)


def test_embedded_image_renders_with_its_caption_right_after_its_heading():
    """Regression test for a real reported bug: embedded diagrams (Context
    Diagram, Container Diagram, Deployment Diagram, ...) rendered as
    anonymous pictures with no title, batched together with no relation to
    which heading/section they actually belonged to. A short heading
    immediately before an <img> in the source must render as that specific
    image's bold caption, directly above it — not as a separate paragraph.
    """
    epic = _node(
        "Epic",
        1,
        "Epic",
        content_sections=[
            {"label": "C4", "blocks": _captioned_image("Context Diagram", "https://dev.azure.com/org/proj/ctx.png")}
        ],
        assets=[_asset("b", "ctx-key", source_url="https://dev.azure.com/org/proj/ctx.png")],
    )
    minio = _fake_minio()
    docx_bytes = build_srs_document(_base_context([epic]), minio)
    doc = Document(BytesIO(docx_bytes))

    paragraphs = [p for p in doc.paragraphs if p.text.strip()]
    caption_idx = next(i for i, p in enumerate(paragraphs) if p.text == "Context Diagram")
    assert paragraphs[caption_idx].runs[0].bold is True
    # The caption must not ALSO show up as an ordinary left-aligned body
    # paragraph elsewhere — it's consumed entirely as the image's caption.
    assert sum(1 for p in paragraphs if p.text == "Context Diagram") == 1
    minio.download_bytes.assert_called_once_with("b", "ctx-key")


def test_multiple_embedded_diagrams_each_get_their_own_caption_in_order():
    epic = _node(
        "Epic",
        1,
        "Epic",
        content_sections=[
            {
                "label": "C4",
                "blocks": (
                    _captioned_image("Context Diagram", "ctx.png") + _captioned_image("Container Diagram", "container.png")
                ),
            }
        ],
        assets=[
            _asset("b", "ctx-key", source_url="ctx.png"),
            _asset("b", "container-key", source_url="container.png"),
        ],
    )
    minio = _fake_minio()
    docx_bytes = build_srs_document(_base_context([epic]), minio)
    doc = Document(BytesIO(docx_bytes))

    caption_texts = [p.text for p in doc.paragraphs if p.text in ("Context Diagram", "Container Diagram")]
    assert caption_texts == ["Context Diagram", "Container Diagram"]
    assert minio.download_bytes.call_count == 2


def test_unmatched_formal_attachment_still_renders_via_fallback():
    """A formal "AttachedFile" attachment never appears inline as an <img>
    tag anywhere in a field's HTML — it must still render, just via the
    existing end-of-node fallback, exactly as it did before this feature.
    """
    epic = _node(
        "Epic",
        1,
        "Epic",
        description_blocks=_text_blocks("Some narrative text."),
        assets=[_asset("b", "attachment-key", source_url=None)],
    )
    minio = _fake_minio()
    build_srs_document(_base_context([epic]), minio)

    minio.download_bytes.assert_called_once_with("b", "attachment-key")


def test_image_with_no_matching_asset_is_skipped_without_crashing():
    """An <img> placeholder that never got downloaded (or download failed)
    must not crash generation — it's simply absent, not a broken document.
    """
    epic = _node(
        "Epic",
        1,
        "Epic",
        content_sections=[{"label": "C4", "blocks": _captioned_image("Context Diagram", "https://example.com/missing.png")}],
        assets=[],
    )
    docx_bytes = build_srs_document(_base_context([epic]), _fake_minio())  # must not raise
    doc = Document(BytesIO(docx_bytes))
    assert "Epic" in "\n".join(p.text for p in doc.paragraphs)

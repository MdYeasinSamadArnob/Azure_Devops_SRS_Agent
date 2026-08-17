"""Increment 4 integration test: runs the full GENERATE chain (normalize ->
run_llm_rules -> render_docx -> convert_pdf -> verify_document) against a
snapshot sealed by the Increment 2/3 fixtures, using the real
Postgres/MinIO/LibreOffice in this stack — no Azure calls are involved in
GENERATE, so nothing needs mocking except the LLM call (kept fast/offline
here; test_llm_adapter.py covers the real Ollama integration separately).
"""

import re
from datetime import datetime, timezone
from io import BytesIO
from unittest.mock import patch

from docx import Document as DocxDocument
from srs_core.db.models import Asset, GeneratedDocument, GenerationJob, JobEvent

from src.tasks import generate_pipeline, import_pipeline


def _cleanup_generation(db_session, generation_job_id):
    asset_ids = [
        doc.asset_id
        for doc in db_session.query(GeneratedDocument).filter(GeneratedDocument.generation_job_id == generation_job_id)
    ]
    db_session.query(GeneratedDocument).filter(GeneratedDocument.generation_job_id == generation_job_id).delete()
    if asset_ids:
        db_session.query(Asset).filter(Asset.id.in_(asset_ids)).delete(synchronize_session=False)
    db_session.query(GenerationJob).filter(GenerationJob.id == generation_job_id).delete()
    db_session.query(JobEvent).filter(JobEvent.job_id == generation_job_id).delete()
    db_session.commit()


def test_generate_chain_produces_docx_and_pdf(db_session, default_tenant, fixture_import_job, azure_mocks):
    import_job, snapshot, connection = fixture_import_job
    import_pipeline.run_import(str(import_job.id), str(snapshot.id), str(connection.id))
    import_pipeline.run_seal(str(import_job.id), str(snapshot.id), str(connection.id))
    db_session.expire_all()

    generation_job = GenerationJob(
        tenant_id=default_tenant.id,
        srs_project_id=snapshot.srs_project_id,
        snapshot_id=snapshot.id,
        status="queued",
        stage="queued",
        params={"formats": ["docx", "pdf"]},
    )
    db_session.add(generation_job)
    db_session.commit()

    try:
        context = generate_pipeline.normalize_content(str(generation_job.id), str(snapshot.id), ["docx", "pdf"])
        # Fixture data is Epic 101 -> Feature 102 -> Story 103: one root, nested.
        assert len(context["roots"]) == 1
        assert context["roots"][0]["work_item_type"] == "Epic"
        assert context["roots"][0]["children"][0]["work_item_type"] == "Feature"
        assert context["total_count"] == 3

        with patch("src.tasks.generate_pipeline.get_default_adapter", return_value=None):
            context = generate_pipeline.run_llm_rules(context)
        assert context["ai_introduction"] is None  # no LLM configured in this test — deterministic fallback used

        context = generate_pipeline.render_docx(context)
        assert context["docx_object_key"]
        assert "roots" not in context

        context = generate_pipeline.convert_pdf(context)
        context = generate_pipeline.verify_document(context)

        db_session.expire_all()
        refreshed_job = db_session.get(GenerationJob, generation_job.id)
        assert refreshed_job.status == "completed"
        assert refreshed_job.stage == "completed"

        docs = (
            db_session.query(GeneratedDocument)
            .filter(GeneratedDocument.generation_job_id == generation_job.id)
            .all()
        )
        formats = {doc.format for doc in docs}
        assert formats == {"docx", "pdf"}
        assert len({doc.generation_id for doc in docs}) == 1  # same generation_id for docx+pdf pair

        from srs_core.storage.minio_client import MinioClient, MinioSettings
        from src.config import get_worker_settings

        settings = get_worker_settings()
        minio = MinioClient(
            MinioSettings(
                endpoint=settings.minio_endpoint,
                access_key=settings.minio_access_key,
                secret_key=settings.minio_secret_key,
                secure=settings.minio_secure,
            )
        )
        # No document_metadata was supplied for this job, so the filename
        # falls back to the root Epic's title ("Epic One" per the fixture),
        # sanitized, plus a YYYYMMDD_HHMMSS timestamp shared by the docx/pdf pair.
        filename_pattern = re.compile(r"^Epic_One_(\d{8}_\d{6})\.(docx|pdf)$")
        docx_doc_asset = None
        timestamps_by_format = {}
        for doc in docs:
            asset_row = db_session.get(Asset, doc.asset_id)
            assert minio.object_exists(asset_row.bucket, asset_row.object_key)
            assert asset_row.byte_size > 0
            match = filename_pattern.match(asset_row.original_filename)
            assert match, f"unexpected filename: {asset_row.original_filename}"
            assert match.group(2) == doc.format
            timestamps_by_format[doc.format] = match.group(1)
            if doc.format == "docx":
                docx_doc_asset = asset_row
        assert timestamps_by_format["docx"] == timestamps_by_format["pdf"]

        # Verify the actual hierarchical structure made it into the document.
        docx_bytes = minio.download_bytes(docx_doc_asset.bucket, docx_doc_asset.object_key)
        doc_obj = DocxDocument(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc_obj.paragraphs)
        assert "Epic One" in full_text
        assert "Feature Two" in full_text
        assert "Traceability Matrix" in full_text
    finally:
        _cleanup_generation(db_session, generation_job.id)


def test_generate_chain_uses_ai_introduction_when_llm_available(db_session, default_tenant, fixture_import_job, azure_mocks):
    """The LLM stage is optional and additive — when configured, its output
    is used; when it fails or isn't configured, the deterministic fallback
    keeps generation working. This exercises the "available" path with a
    mocked adapter so the suite doesn't depend on network access to a real
    LLM (see test_llm_adapter.py's live test for that).
    """
    import_job, snapshot, connection = fixture_import_job
    import_pipeline.run_import(str(import_job.id), str(snapshot.id), str(connection.id))
    import_pipeline.run_seal(str(import_job.id), str(snapshot.id), str(connection.id))
    db_session.expire_all()

    generation_job = GenerationJob(
        tenant_id=default_tenant.id,
        srs_project_id=snapshot.srs_project_id,
        snapshot_id=snapshot.id,
        status="queued",
        stage="queued",
        params={"formats": ["docx"]},
    )
    db_session.add(generation_job)
    db_session.commit()

    try:
        context = generate_pipeline.normalize_content(str(generation_job.id), str(snapshot.id), ["docx"])

        fake_adapter = type("FakeAdapter", (), {"complete": lambda self, *a, **kw: "This system manages customer data."})()
        with patch("src.tasks.generate_pipeline.get_default_adapter", return_value=fake_adapter):
            context = generate_pipeline.run_llm_rules(context)
        assert context["ai_introduction"] == "This system manages customer data."

        context = generate_pipeline.render_docx(context)

        from srs_core.storage.minio_client import MinioClient, MinioSettings
        from src.config import get_worker_settings

        settings = get_worker_settings()
        minio = MinioClient(
            MinioSettings(
                endpoint=settings.minio_endpoint,
                access_key=settings.minio_access_key,
                secret_key=settings.minio_secret_key,
                secure=settings.minio_secure,
            )
        )
        docx_bytes = minio.download_bytes(context["docx_bucket"], context["docx_object_key"])
        doc_obj = DocxDocument(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc_obj.paragraphs)
        assert "This system manages customer data." in full_text
        assert "AI-generated summary" in full_text
    finally:
        _cleanup_generation(db_session, generation_job.id)


def test_render_docx_survives_malformed_html_and_deeply_nested_hierarchy(db_session, default_tenant, fixture_import_job):
    """Regression test for a real production bug: a work item description
    containing an unclosed HTML tag (very common in real Azure rich-text
    fields) used to silently corrupt the old docxtpl-rendered DOCX's XML
    from that item onward, so everything after it vanished without error.
    The engine no longer uses docxtpl for the body at all (python-docx's
    own text APIs escape correctly by construction) — this proves the new
    engine handles it, and that Epic > Feature > Task nesting with
    acceptance criteria renders correctly.
    """
    _import_job, snapshot, _connection = fixture_import_job

    generation_job = GenerationJob(
        tenant_id=default_tenant.id,
        srs_project_id=snapshot.srs_project_id,
        snapshot_id=snapshot.id,
        status="queued",
        stage="queued",
        params={"formats": ["docx"]},
    )
    db_session.add(generation_job)
    db_session.commit()

    context = {
        "generation_job_id": str(generation_job.id),
        "snapshot_id": str(snapshot.id),
        "tenant_id": str(default_tenant.id),
        "srs_project_id": str(snapshot.srs_project_id),
        "source_url": "https://dev.azure.com/org/proj",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "formats": ["docx"],
        "total_count": 3,
        "ai_introduction": None,
        "roots": [
            {
                "work_item_type": "Epic",
                "azure_work_item_id": 1,
                "title": "First item",
                "state": "New",
                "description_blocks": [],
                "acceptance_criteria_blocks": [],
                "azure_url": "#1",
                "assets": [],
                "children": [
                    {
                        "work_item_type": "Feature",
                        "azure_work_item_id": 2,
                        "title": "Poisoned item",
                        "state": "New",
                        # Real-world Azure HTML: an unclosed tag used to corrupt raw XML insertion.
                        "description_blocks": [
                            {
                                "type": "text",
                                "runs": [
                                    {
                                        "text": "<div>some text <unclosed tag here",
                                        "bold": False,
                                        "italic": False,
                                        "underline": False,
                                    }
                                ],
                            }
                        ],
                        "acceptance_criteria_blocks": [
                            {
                                "type": "list",
                                "items": [
                                    [{"text": "must handle X", "bold": False, "italic": False, "underline": False}],
                                    [{"text": "must handle Y", "bold": False, "italic": False, "underline": False}],
                                ],
                            }
                        ],
                        "azure_url": "#2",
                        "assets": [],
                        "children": [
                            {
                                "work_item_type": "Task",
                                "azure_work_item_id": 3,
                                "title": "Third item after the poisoned one",
                                "state": "New",
                                "description_blocks": [
                                    {
                                        "type": "text",
                                        "runs": [
                                            {"text": "plain description", "bold": False, "italic": False, "underline": False}
                                        ],
                                    }
                                ],
                                "acceptance_criteria_blocks": [],
                                "azure_url": "#3",
                                "assets": [],
                                "children": [],
                            }
                        ],
                    }
                ],
            }
        ],
    }

    try:
        result = generate_pipeline.render_docx(context)

        from srs_core.storage.minio_client import MinioClient, MinioSettings
        from src.config import get_worker_settings

        settings = get_worker_settings()
        minio = MinioClient(
            MinioSettings(
                endpoint=settings.minio_endpoint,
                access_key=settings.minio_access_key,
                secret_key=settings.minio_secret_key,
                secure=settings.minio_secure,
            )
        )
        docx_bytes = minio.download_bytes(result["docx_bucket"], result["docx_object_key"])
        doc = DocxDocument(BytesIO(docx_bytes))
        full_text = "\n".join(p.text for p in doc.paragraphs)

        assert "First item" in full_text
        assert "Poisoned item" in full_text
        assert "some text <unclosed tag here" in full_text  # rendered as literal safe text, not corrupted XML
        assert "Third item after the poisoned one" in full_text
        assert "must handle X" in full_text
        assert "must handle Y" in full_text
    finally:
        _cleanup_generation(db_session, generation_job.id)
